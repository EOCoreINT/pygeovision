"""Excel workbook (openpyxl): summary, tables with live formulas, native charts, maps."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pygeovision.reporting import fmt as F
from pygeovision.reporting.model import Report, Table
from pygeovision.reporting.palettes import ACCENT, ACCENT_SOFT, INK, MUTED, WARN_SOFT

_FONT = "Arial"
_SHEET_ORDER = ("Summary", "Metrics", "Classes", "Statistics", "Timeseries", "Inputs", "Vectors", "Files", "Parameters", "Maps")


def _argb(hex_color: str) -> str:
    return "FF" + hex_color.lstrip("#").upper()


def write_xlsx(report: Report, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True

    f_base = Font(name=_FONT, size=10, color=_argb(INK))
    f_bold = Font(name=_FONT, size=10, bold=True, color=_argb(INK))
    f_head = Font(name=_FONT, size=10, bold=True, color="FFFFFFFF")
    f_title = Font(name=_FONT, size=18, bold=True, color=_argb(ACCENT))
    f_h2 = Font(name=_FONT, size=12, bold=True, color=_argb(ACCENT))
    f_muted = Font(name=_FONT, size=9, color=_argb(MUTED))
    fill_head = PatternFill("solid", fgColor=_argb(ACCENT))
    fill_soft = PatternFill("solid", fgColor=_argb(ACCENT_SOFT))
    fill_warn = PatternFill("solid", fgColor=_argb(WARN_SOFT))
    thin = Side(style="thin", color="FFD5DBE1")
    border = Border(bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    sheets: dict[str, Any] = {}

    def sheet(name: str):
        if name not in sheets:
            ws = wb.create_sheet(name[:31])
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.sheet_view.showGridLines = False
            sheets[name] = ws
        return sheets[name]

    def put(ws, row: int, col: int, value: Any, *, font=None, fill=None, fmt: str | None = None, align=None, b=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = font or f_base
        if fill is not None:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        if align is not None:
            c.alignment = align
        if b is not None:
            c.border = b
        return c

    def widths(ws, spec: dict[int, float]):
        for col, w in spec.items():
            cur = ws.column_dimensions[get_column_letter(col)].width
            ws.column_dimensions[get_column_letter(col)].width = max(w, cur if cur and cur != 13 else 0)

    # ---- Summary --------------------------------------------------------
    ws = sheet("Summary")
    put(ws, 1, 1, report.title, font=f_title)
    ws.row_dimensions[1].height = 28
    put(ws, 2, 1, report.subtitle, font=f_muted)
    r = 4
    if report.kpis:
        for i, (lab, val, _) in enumerate(report.kpis[:4]):
            col = 1 + i * 2
            put(ws, r, col, val, font=Font(name=_FONT, size=15, bold=True, color=_argb(ACCENT)), fill=fill_soft)
            put(ws, r + 1, col, lab, font=f_muted, fill=fill_soft, align=wrap)
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
        ws.row_dimensions[r].height = 26
        ws.row_dimensions[r + 1].height = 30
        r += 3
    for s in report.sections:
        if s.title == "Executive summary":
            put(ws, r, 1, "Executive summary", font=f_h2)
            r += 1
            for b in s.blocks:
                lines = b.items if b.kind == "bullets" else [b.text] if b.kind in ("p", "callout") else []
                for line in lines:
                    warn = b.kind == "callout" and b.level == "warn"
                    put(ws, r, 1, ("• " if b.kind == "bullets" else "") + str(line), fill=fill_warn if warn else None, align=wrap)
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                    ws.row_dimensions[r].height = max(16, 15 * (1 + len(str(line)) // 105))
                    r += 1
            r += 1
        elif s.title == "Study area and inputs":
            put(ws, r, 1, "Study area", font=f_h2)
            r += 1
            for b in s.blocks:
                if b.kind == "kv":
                    for k_, v_ in b.items:
                        put(ws, r, 1, k_, font=f_muted, b=border)
                        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
                        put(ws, r, 4, str(v_), b=border)
                        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
                        r += 1
            r += 1
    if report.cover_image and Path(report.cover_image).exists():
        put(ws, r, 1, "Map", font=f_h2)
        img = XLImage(str(report.cover_image))
        ratio = img.height / img.width
        img.width, img.height = 880, int(880 * ratio)
        ws.add_image(img, f"A{r + 1}")
    widths(ws, {i: 17 for i in range(1, 9)})

    # ---- tables ---------------------------------------------------------
    cursor: dict[str, int] = {}

    def start_row(name: str) -> int:
        return cursor.get(name, 1)

    def write_table(name: str, t: Table, chart: bool = False) -> None:
        w = sheet(name)
        row = start_row(name)
        put(w, row, 1, t.title, font=f_h2)
        row += 1
        has_sw = t.colors is not None
        off = 2 if has_sw else 1       # first data column
        for j, c in enumerate(t.cols):
            put(w, row, off + j, c.name, font=f_head, fill=fill_head,
                align=Alignment(horizontal="right" if c.align == "right" else "left", vertical="center", wrap_text=True))
        if has_sw:
            put(w, row, 1, "", fill=fill_head)
        w.row_dimensions[row].height = 30
        hdr = row
        row += 1
        first = row
        for i, vals in enumerate(t.rows):
            for j, (c, v) in enumerate(zip(t.cols, vals)):
                num = c.fmt != "text" and not F.is_missing(v) and isinstance(v, (int, float))
                put(w, row, off + j, v if num else (F.NA if v is None else str(v)) if c.fmt == "text" else F.NA,
                    fmt=F.EXCEL_FORMATS.get(c.fmt) if num else None, b=border,
                    align=Alignment(horizontal="right" if c.align == "right" else "left", wrap_text=True, vertical="top"))
            if has_sw and t.colors and t.colors[i]:
                put(w, row, 1, "", fill=PatternFill("solid", fgColor=_argb(t.colors[i])), b=border)
            row += 1
        last = row - 1
        names = [c.name for c in t.cols]
        if t.total_row and t.rows:
            put(w, row, off, "Total", font=f_bold)
            for j in t.sum_cols:
                col = get_column_letter(off + j)
                put(w, row, off + j, f"=SUM({col}{first}:{col}{last})", font=f_bold, fmt=F.EXCEL_FORMATS.get(t.cols[j].fmt))
            # share of valid area as a live formula on the total of the area column
            if "Area (ha)" in names and "Share of valid area" in names:
                a = get_column_letter(off + names.index("Area (ha)"))
                s_col = off + names.index("Share of valid area")
                for rr in range(first, last + 1):
                    put(w, rr, s_col, f"=IF(SUM(${a}${first}:${a}${last})=0,0,{a}{rr}/SUM(${a}${first}:${a}${last}))", fmt="0.0%", b=border,
                        align=Alignment(horizontal="right"))
                put(w, row, s_col, f"=SUM({get_column_letter(s_col)}{first}:{get_column_letter(s_col)}{last})", font=f_bold, fmt="0.0%")
            row += 1
        if t.note:
            put(w, row, 1, t.note, font=f_muted)
            row += 1
        widths(w, {1: 4 if has_sw else 34})
        for j, c in enumerate(t.cols):
            widths(w, {off + j: 14 + 12 * (c.weight >= 2.5)})
        if chart and t.rows and "Area (ha)" in names:
            ch = BarChart()
            ch.type = "bar"
            ch.title = t.title
            ch.style = 10
            ch.height, ch.width = 7.5, 15
            data = Reference(w, min_col=off + names.index("Area (ha)"), min_row=hdr, max_row=last)
            cats = Reference(w, min_col=off, min_row=first, max_row=last)
            ch.add_data(data, titles_from_data=True)
            ch.set_categories(cats)
            ch.legend = None
            ch.y_axis.title = "Area (ha)"
            w.add_chart(ch, f"{get_column_letter(off + len(t.cols) + 1)}{hdr - 1}")
            row = max(row, hdr + 17)
        cursor[name] = row + 2

    for s in report.sections:
        for b in s.blocks:
            if b.kind == "table" and b.table is not None and b.table.sheet:
                t = b.table
                if t.raw is not None:           # (label, value, code, note) rows keep numeric cells
                    w = sheet(t.sheet)
                    row = start_row(t.sheet)
                    put(w, row, 1, t.title, font=f_h2)
                    row += 1
                    for j, name in enumerate(("Measure", "Value", "Note")):
                        put(w, row, 1 + j, name, font=f_head, fill=fill_head, align=Alignment(horizontal="right" if j == 1 else "left"))
                    row += 1
                    for label, val, code, note in t.raw:
                        put(w, row, 1, label, b=border, align=wrap)
                        num = code not in ("text", "bool") and not F.is_missing(val) and isinstance(val, (int, float)) and not isinstance(val, bool)
                        put(w, row, 2, val if num else F.fmt_value(val, code), fmt=F.EXCEL_FORMATS.get(code) if num else None, b=border,
                            align=Alignment(horizontal="right"))
                        put(w, row, 3, note, font=f_muted, b=border, align=wrap)
                        row += 1
                    widths(w, {1: 42, 2: 18, 3: 52})
                    cursor[t.sheet] = row + 2
                else:
                    write_table(t.sheet, t, chart=(t.sheet == "Classes"))

    # time-series chart
    if report.timeseries and "Timeseries" in sheets:
        w = sheets["Timeseries"]
        n = len(report.timeseries)
        ch = LineChart()
        ch.title = "Scene means by date"
        ch.height, ch.width = 8, 16
        ch.style = 12
        data = Reference(w, min_col=2, max_col=4, min_row=2, max_row=2 + n)
        cats = Reference(w, min_col=1, min_row=3, max_row=2 + n)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        w.add_chart(ch, "H2")

    # parameters
    for s in report.sections:
        if s.title == "Method and parameters":
            w = sheet("Parameters")
            row = 1
            put(w, row, 1, "Method", font=f_h2)
            row += 1
            for b in s.blocks:
                if b.kind == "p":
                    put(w, row, 1, b.text, align=wrap)
                    w.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
                    w.row_dimensions[row].height = max(16, 15 * (1 + len(b.text) // 100))
                    row += 1
                elif b.kind == "kv":
                    row += 1
                    put(w, row, 1, "Run parameters", font=f_h2)
                    row += 1
                    for k_, v_ in b.items:
                        put(w, row, 1, k_, font=f_muted, b=border)
                        put(w, row, 2, str(v_), b=border, align=wrap)
                        w.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
                        row += 1
            widths(w, {1: 34, 2: 40, 3: 40})

    # feature attribute tables
    for name, rows in report.feature_tables.items():
        w = sheet(f"Features {name}"[:31])
        total = report.feature_totals.get(name, len(rows))
        put(w, 1, 1, f"{name}: {total:,} polygons", font=f_h2)
        note = (f"Showing the {len(rows):,} largest of {total:,}. The full set is in the GeoPackage/GeoJSON/Shapefile."
                if total > len(rows) else "All polygons.")
        put(w, 2, 1, note, font=f_muted)
        cols = [("patch_id", "Patch", "0"), ("class_id", "Class value", "0"), ("class_name", "Class", "@"), ("area_ha", "Area (ha)", "#,##0.00"),
                ("pct_aoi", "% of AOI", "0.000"), ("cx", "Centre lon", "0.00000"), ("cy", "Centre lat", "0.00000")]
        for j, (_, head, _) in enumerate(cols):
            put(w, 4, 1 + j, head, font=f_head, fill=fill_head)
        for i, rec in enumerate(rows):
            for j, (k, _, nf) in enumerate(cols):
                put(w, 5 + i, 1 + j, rec.get(k), fmt=nf)
        w.freeze_panes = "A5"
        w.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(rows)}"
        widths(w, {1: 10, 2: 12, 3: 30, 4: 14, 5: 12, 6: 14, 7: 14})

    # maps
    figs = [b.figure for s in report.sections for b in s.blocks if b.kind == "figure" and b.figure is not None and b.figure.path.exists()]
    if figs:
        w = sheet("Maps")
        row = 1
        for fg in figs:
            put(w, row, 1, fg.caption, font=f_h2)
            img = XLImage(str(fg.path))
            ratio = img.height / img.width
            img.width, img.height = 900, int(900 * ratio)
            w.add_image(img, f"A{row + 1}")
            row += int(img.height / 20) + 4

    # warnings
    if report.warnings:
        w = sheet("Parameters")
        row = w.max_row + 2
        put(w, row, 1, "Notes on this run", font=f_h2)
        for i, msg in enumerate(report.warnings, 1):
            put(w, row + i, 1, msg, align=wrap)
            w.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=3)
            w.row_dimensions[row + i].height = 30

    for w_ in sheets.values():                      # printing: landscape, one page wide
        w_.page_setup.orientation = "landscape"
        w_.page_setup.fitToWidth = 1
        w_.page_setup.fitToHeight = 0
        w_.sheet_properties.pageSetUpPr.fitToPage = True
    ordered = [n for n in _SHEET_ORDER if n in sheets] + [n for n in sheets if n not in _SHEET_ORDER]
    wb._sheets = [wb[sheets[n].title] for n in ordered]
    wb.properties.title = report.title
    wb.properties.creator = "PyGeoVision"
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out
