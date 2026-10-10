"""Word report (python-docx)."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pygeovision.reporting import fmt as F
from pygeovision.reporting.model import Block, Report, Table
from pygeovision.reporting.palettes import ACCENT, ACCENT_SOFT, INK, MUTED, WARN, WARN_SOFT

_FONT = "Arial"
_TEXT_WIDTH_CM = 17.0


def _rgb(hex_color: str):
    from docx.shared import RGBColor
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    tc_pr.append(shd)


def _cell_borders(cell: Any, **edges: tuple[str, int, str]) -> None:
    """edges: left/right/top/bottom -> (style, size in 1/8 pt, colour)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, (style, size, color) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color.lstrip("#"))
        borders.append(el)


def _cell_margins(table: Any, top: int = 50, bottom: int = 50, left: int = 90, right: int = 90) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _repeat_header(row: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _no_split(row: Any) -> None:
    from docx.oxml import OxmlElement
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _field(run: Any, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False, size: float = 8.5, color: str = INK, align: str = "left") -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    cell.text = ""
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        par.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        run = par.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)


def write_docx(report: Report, path: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before in (("Heading 1", 16, ACCENT, 14), ("Heading 2", 11.5, INK, 9)):
        st = doc.styles[name]
        st.font.name = _FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = _rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    for name in ("Heading 1", "Heading 2", "List Bullet"):
        rpr = doc.styles[name].element.get_or_add_rPr()
        rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if rfonts is not None:
            for attr in list(rfonts.attrib):
                if "theme" in attr.lower():
                    del rfonts.attrib[attr]

    # header and footer
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(f"{report.title}    |    PyGeoVision {report.version}")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUTED)
    foot_style = doc.styles["Footer"]
    foot_style.font.size, foot_style.font.bold = Pt(8), False
    foot_style.font.color.rgb = _rgb(MUTED)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(f"Generated {report.generated}    |    Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUTED)
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    r2.font.color.rgb = _rgb(MUTED)
    _field(r2, "PAGE")
    r3 = fp.add_run(" of ")
    r3.font.size = Pt(8)
    r3.font.color.rgb = _rgb(MUTED)
    r4 = fp.add_run()
    r4.font.size = Pt(8)
    r4.font.color.rgb = _rgb(MUTED)
    _field(r4, "NUMPAGES")

    # cover band
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    cell = band.rows[0].cells[0]
    cell.width = Cm(_TEXT_WIDTH_CM)
    _shade(cell, ACCENT)
    _cell_margins(band, 220, 220, 260, 260)
    band.columns[0].width = Cm(_TEXT_WIDTH_CM)
    cell.text = ""
    p0 = cell.paragraphs[0]
    k = p0.add_run("PYGEOVISION  ·  GEOSPATIAL ANALYSIS REPORT")
    k.font.size, k.font.bold = Pt(8.5), True
    k.font.color.rgb = _rgb("#8fb8d3")
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(10)
    t = p1.add_run(report.title)
    t.font.size, t.font.bold = Pt(24), True
    t.font.color.rgb = _rgb("#ffffff")
    p2 = cell.add_paragraph()
    s = p2.add_run(report.subtitle)
    s.font.size = Pt(11)
    s.font.color.rgb = _rgb("#dbe9f3")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    if report.kpis:
        kt = doc.add_table(rows=1, cols=min(4, len(report.kpis)))
        kt.autofit = False
        w = _TEXT_WIDTH_CM / len(kt.columns)
        for i, (lab, val, _) in enumerate(report.kpis[:4]):
            c = kt.rows[0].cells[i]
            c.width = Cm(w)
            _shade(c, ACCENT_SOFT)
            _cell_borders(c, left=("single", 18, ACCENT))
            c.text = ""
            pv = c.paragraphs[0]
            pv.paragraph_format.space_after = Pt(1)
            rv = pv.add_run(val)
            rv.bold, rv.font.size = True, Pt(15)
            rv.font.color.rgb = _rgb(ACCENT)
            pl = c.add_paragraph()
            pl.paragraph_format.space_after = Pt(0)
            rl = pl.add_run(lab)
            rl.font.size = Pt(8)
            rl.font.color.rgb = _rgb(MUTED)
        _cell_margins(kt, 90, 90, 120, 100)
        for i in range(len(kt.columns)):
            kt.columns[i].width = Cm(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_picture(p: Path, caption: str, width_cm: float = _TEXT_WIDTH_CM):
        par = doc.add_paragraph()
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(str(p), width=Cm(width_cm))
        cp = doc.add_paragraph()
        r = cp.add_run(caption)
        r.italic, r.font.size = True, Pt(8)
        r.font.color.rgb = _rgb(MUTED)


    def add_table(t: Table) -> None:
        h = doc.add_paragraph()
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(6)
        hr = h.add_run(t.title)
        hr.bold, hr.font.size = True, Pt(10)
        has_sw = t.colors is not None
        ncols = len(t.cols) + (1 if has_sw else 0)
        tbl = doc.add_table(rows=1, cols=ncols)
        tbl.autofit = False
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        sw_cm = 0.45 if has_sw else 0.0
        total_w = sum(c.weight for c in t.cols)
        widths = ([sw_cm] if has_sw else []) + [(_TEXT_WIDTH_CM - sw_cm) * c.weight / total_w for c in t.cols]
        hdr = tbl.rows[0]
        _repeat_header(hdr)
        names = ([""] if has_sw else []) + [c.name for c in t.cols]
        aligns = (["left"] if has_sw else []) + [c.align for c in t.cols]
        for i, (n, a) in enumerate(zip(names, aligns)):
            hdr.cells[i].width = Cm(widths[i])
            _shade(hdr.cells[i], ACCENT)
            _set_cell_text(hdr.cells[i], n, bold=True, color="#ffffff", align=a)
        for ri, row in enumerate(t.rows):
            cells = tbl.add_row()
            _no_split(cells)
            vals = ([""] if has_sw else []) + [
                (F.fmt_value(v, c.fmt) if c.fmt != "text" else (F.NA if v is None else str(v))) for c, v in zip(t.cols, row)]
            for i, v in enumerate(vals):
                cells.cells[i].width = Cm(widths[i])
                _set_cell_text(cells.cells[i], v, align=aligns[i])
                if ri % 2 == 1:
                    _shade(cells.cells[i], "#f5f7f9")
            if has_sw and t.colors and t.colors[ri]:
                _shade(cells.cells[0], t.colors[ri])
        _cell_margins(tbl)
        for i, w_ in enumerate(widths):
            tbl.columns[i].width = Cm(w_)
        if len(t.rows) <= 18:                       # keep a short table, its title and its note on one page
            for row_ in tbl.rows:
                for cell_ in row_.cells:
                    for par_ in cell_.paragraphs:
                        par_.paragraph_format.keep_with_next = True
        if t.note:
            n = doc.add_paragraph()
            r = n.add_run(t.note)
            r.italic, r.font.size = True, Pt(8)
            r.font.color.rgb = _rgb(MUTED)
        else:
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_callout(text: str, level: str) -> None:
        bg, bar = (WARN_SOFT, WARN) if level == "warn" else (ACCENT_SOFT, ACCENT)
        tb = doc.add_table(rows=1, cols=1)
        tb.autofit = False
        c = tb.rows[0].cells[0]
        c.width = Cm(_TEXT_WIDTH_CM)
        _shade(c, bg)
        _cell_borders(c, left=("single", 24, bar))
        _set_cell_text(c, text, size=9.5)
        _cell_margins(tb, 80, 80, 160, 120)
        tb.columns[0].width = Cm(_TEXT_WIDTH_CM)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def add_kv(pairs) -> None:
        tb = doc.add_table(rows=0, cols=2)
        tb.autofit = False
        for k_, v_ in pairs:
            row = tb.add_row()
            _no_split(row)
            row.cells[0].width, row.cells[1].width = Cm(5.8), Cm(11.2)
            _set_cell_text(row.cells[0], k_, size=8.5, color=MUTED)
            _set_cell_text(row.cells[1], str(v_), size=9)
            for c in row.cells:
                _cell_borders(c, bottom=("single", 4, "#d5dbe1"))
        _cell_margins(tb, 40, 40, 60, 60)
        tb.columns[0].width, tb.columns[1].width = Cm(5.8), Cm(11.2)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_block(b: Block) -> None:
        if b.kind == "p":
            doc.add_paragraph(b.text)
        elif b.kind == "bullets":
            for it in b.items:
                doc.add_paragraph(str(it), style="List Bullet")
        elif b.kind == "table" and b.table is not None:
            add_table(b.table)
        elif b.kind == "figure" and b.figure is not None and b.figure.path.exists():
            add_picture(b.figure.path, b.figure.caption)
        elif b.kind == "callout":
            add_callout(b.text, b.level)
        elif b.kind == "kv":
            add_kv(b.items)

    first = True
    for s in report.sections:
        if s.page_break and not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        doc.add_heading(s.title, level=1)
        for b in s.blocks:
            add_block(b)
    if report.warnings:
        doc.add_heading("Notes on this run", level=1)
        doc.add_paragraph("Parts of the deliverables that could not be produced, or were changed, and why:")
        for w in report.warnings:
            doc.add_paragraph(w, style="List Bullet")

    doc.core_properties.title = report.title
    doc.core_properties.author = "PyGeoVision"
    doc.core_properties.subject = report.subtitle
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
