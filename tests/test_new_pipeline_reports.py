"""Every new pipeline produces a complete report: maps, charts, tables, vectors, and no warnings about missing pieces.

Each case runs the real pipeline on a tiny hand-made world (offline), then builds the deliverables and checks the
Word text and the vector files.
"""
from __future__ import annotations

import json
from types import SimpleNamespace

import geopandas as gpd
import numpy as np
import pytest
from shapely.geometry import LineString, box

from pygeovision.ai.pipelines.reference_products import (
    ActiveFirePipeline,
    EvapotranspirationPipeline,
    SurfaceWaterDynamicsPipeline,
    UrbanExpansionPipeline,
)
from pygeovision.data import footprints as fp
from pygeovision.reporting import generate_deliverables
from pygeovision.reporting.specs import SPECS
from tests import test_reference_pipelines as rp
from tests import test_terrain_pipelines as tp
from tests.pipeline_synth import BBOX as LONDON
from tests.pipeline_synth import patch_downloads, write_raster
from tests.test_timeseries_pipelines import N, mndwi_bands, ndvi_field, red_nir, tif

BAD_WARNING = ("expected output", "could not", "failed", "not found", "skipped")


def docx_text(path):
    from docx import Document
    d = Document(str(path))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def build(result, bbox, out, formats="docx,xlsx", **kw):
    assert result.success, result.error
    d = generate_deliverables(result, bbox, out, report_formats=formats, **kw)
    bad = [w for w in d.warnings if any(k in w for k in BAD_WARNING)]
    assert not bad, bad
    assert d.maps and d.maps[0].exists() and d.charts, "a map and a chart are expected"
    return d


# ---------------------------------------------------------------------------------------------------------------------
def test_landcover_change_report(tmp_path):
    g, pc, r = rp.TestLandCoverChange()._run(tmp_path)
    d = build(r, rp.BBOX, tmp_path / "out")
    text = docx_text(d.reports["docx"])
    assert "Largest class changes" in text and "Trees" in text and "Built area" in text
    assert "Land cover 2018" in text and "Land cover 2022" in text            # the two year maps are extra layers
    assert "Impact Observatory" in text                                       # the product note callout / reference
    assert len(d.maps) == 3
    import openpyxl
    wb = openpyxl.load_workbook(d.reports["xlsx"])
    assert "Transitions" in wb.sheetnames


def test_urban_expansion_report(tmp_path):
    g = rp.grid_for(10)
    a = np.full((g.height, g.width), 5, np.uint8)
    a[0:20, 0:20] = 7
    b = a.copy()
    b[0:10, 0:20] = 8
    b[50:80, 50:60] = 7
    pc = rp.FakePC({"io-lulc-annual-v02": [
        rp.item("a", rp.write_tile(tmp_path / "a.tif", a, g), "2018-01-01T00:00:00Z", "2019-01-01T00:00:00Z"),
        rp.item("b", rp.write_tile(tmp_path / "b.tif", b, g), "2023-01-01T00:00:00Z", "2024-01-01T00:00:00Z")]})
    p = UrbanExpansionPipeline(None)
    p.pc = pc
    r = p.run(rp.BBOX, tmp_path / "out", date_before=2018, date_after=2023)
    d = build(r, rp.BBOX, tmp_path / "out")
    text = docx_text(d.reports["docx"])
    assert "New built area" in text and "gross gain" in text
    gdf = gpd.read_file(d.vectors["gpkg"][0], layer="built_change")
    assert set(gdf["class_name"]) == {"New built area", "Built area lost"}


def test_forest_change_report(tmp_path):
    r, g = rp.TestForestChange()._run(tmp_path)
    d = build(r, rp.BBOX, tmp_path / "out")
    assert "Forest loss" in docx_text(d.reports["docx"])
    gdf = gpd.read_file(d.vectors["gpkg"][0], layer="forest_change")
    assert set(gdf["class_name"]) == {"Forest loss", "Forest gain"}


def test_surface_water_report(tmp_path):
    g = rp.grid_for(30)
    h, w = g.height, g.width
    tr = np.zeros((h, w), np.uint8)
    for i in range(10):
        tr[10 * i:10 * i + 10, 0:10] = i + 1
    occ = np.where(tr > 0, 80, 0).astype(np.uint8)
    sea = np.where(tr == 4, 12, 0).astype(np.uint8)
    files = {n: rp.write_tile(tmp_path / f"{n}.tif", a, g, nodata=None) for n, a in
             (("transitions", tr), ("occurrence", occ), ("seasonality", sea))}
    it = {"id": "0E_10Nv1_3_2020", "properties": {"datetime": "2020-07-01T00:00:00Z"}, "assets": {k: {"href": v} for k, v in files.items()}}
    p = SurfaceWaterDynamicsPipeline(None)
    p.pc = rp.FakePC({"jrc-gsw": [it]})
    r = p.run(rp.BBOX, tmp_path / "out")
    d = build(r, rp.BBOX, tmp_path / "out")
    assert len(d.maps) == 3 and "Lost permanent" in docx_text(d.reports["docx"])


def test_evapotranspiration_report(tmp_path):
    g = rp.grid_for(500, rp.COARSE)
    et = np.full((g.height, g.width), 12345, np.uint16)
    pet = np.full((g.height, g.width), 20000, np.uint16)
    et[5:10, 5:10] = 8000
    t = rp.TestEvapotranspiration()
    p = EvapotranspirationPipeline(None)
    p.pc = rp.FakePC({"modis-16A3GF-061": t._items(tmp_path, et, pet, g)})
    r = p.run(rp.COARSE, tmp_path / "out", date="2022-06", resolution_m=500)
    d = build(r, rp.COARSE, tmp_path / "out")
    assert "million m" in docx_text(d.reports["docx"])


def test_active_fire_report_has_the_detection_points_and_a_per_day_table(tmp_path):
    g, pc = rp.TestActiveFire()._world(tmp_path)
    p = ActiveFirePipeline(None)
    p.pc = pc
    r = p.run(rp.COARSE, tmp_path / "out", date="2024-02", min_confidence="nominal")
    d = build(r, rp.COARSE, tmp_path / "out", vector_formats="gpkg")
    text = docx_text(d.reports["docx"])
    assert "Detections per day" in text and "2024-02-11" in text
    pts = gpd.read_file(d.vectors["gpkg"][0], layer="fire_detections")
    assert len(pts) == 3 and {"confidence", "frp_mw"} <= set(pts.columns)


def test_terrain_report(tmp_path):
    r, grid = tp.TestTerrainAnalysis()._run(tmp_path, lambda rr, cc: 3.0 * cc + 0.5 * rr)
    d = build(r, tp.BBOX, tmp_path / "out")
    assert len(d.maps) == 5 and "Hillshade" in docx_text(d.reports["docx"])


def test_drainage_report_exports_the_stream_lines_unchanged(tmp_path):
    r, grid, c0 = tp.TestDrainageAnalysis()._valley(tmp_path)
    d = build(r, tp.BBOX, tmp_path / "out", vector_formats="gpkg,geojson")
    text = docx_text(d.reports["docx"])
    assert "Height above nearest drainage" in text and "Strahler" in text
    streams = gpd.read_file(d.vectors["gpkg"][0], layer="streams")
    assert list(streams["strahler"]) == [1] and streams.geom_type.iloc[0] in ("LineString", "MultiLineString")
    assert "area_ha" not in streams.columns                 # no area for a line
    assert any(p.name.endswith("_streams.geojson") for p in d.vectors["geojson"])


def deg(x):
    return x * 0.0001


def test_building_footprints_report_keeps_the_polygons(tmp_path, monkeypatch):
    bbox = (0.45, 5.90, 0.47, 5.92)
    geoms = [box(0.4601 + i * 0.0003, 5.9101, 0.4601 + i * 0.0003 + deg(1), 5.9101 + deg(1)) for i in range(20)]
    monkeypatch.setattr(fp, "microsoft_footprints", lambda b, **k: gpd.GeoDataFrame({"id": range(20)}, geometry=geoms, crs="EPSG:4326"))
    from pygeovision.ai.pipelines import BuildingFootprintsPipeline
    r = BuildingFootprintsPipeline(None).run(bbox, tmp_path / "out", source="microsoft")
    d = build(r, bbox, tmp_path / "out", formats="docx,xlsx,pdf")
    text = docx_text(d.reports["docx"])
    assert "Microsoft" in text and "20" in text
    assert "accuracy is not assessed" not in text           # published dataset: no network was involved
    layer = gpd.read_file(d.vectors["gpkg"][0], layer="buildings")
    assert len(layer) == 20 and layer["area_ha"].gt(0).all() and (layer["source"] == "microsoft").all()
    assert d.reports["pdf"].stat().st_size > 5000
    import openpyxl
    assert any(n.startswith("Features buildings") for n in openpyxl.load_workbook(d.reports["xlsx"]).sheetnames)


def test_relative_output_folder_still_yields_one_geojson_feature_per_building(tmp_path, monkeypatch):
    """`pygeovision channel ... -o out/bld` (a relative folder) once lost the buildings layer: the path was joined twice."""
    bbox = (0.45, 5.90, 0.47, 5.92)
    geoms = [box(0.4601 + i * 0.0003, 5.9101, 0.4601 + i * 0.0003 + deg(1), 5.9101 + deg(1)) for i in range(20)]
    monkeypatch.setattr(fp, "microsoft_footprints", lambda b, **k: gpd.GeoDataFrame({"id": range(20)}, geometry=geoms, crs="EPSG:4326"))
    monkeypatch.chdir(tmp_path)
    from pygeovision.ai.pipelines import BuildingFootprintsPipeline
    r = BuildingFootprintsPipeline(None).run(bbox, "out/bld", source="microsoft")
    d = build(r, bbox, "out/bld", formats="docx", vector_formats="gpkg,geojson,shp")
    gj = [p for p in d.vectors["geojson"] if p.name.endswith("_buildings.geojson")]
    assert gj, "the buildings GeoJSON is missing"
    g = gpd.read_file(gj[0])
    assert len(g) == 20 and g.geom_type.eq("Polygon").all()          # separate shapes, not one merged geometry
    assert g.geometry.apply(lambda x: x.area).gt(0).all() and g["patch_id"].is_unique
    assert not any("could not be read" in w for w in d.warnings)
    assert len(gpd.read_file(d.vectors["gpkg"][0], layer="buildings")) == 20
    assert any(p.name.endswith("_buildings.shp") for p in d.vectors["shp"])


def test_building_footprints_with_a_network_still_says_accuracy_is_not_assessed(tmp_path):
    out = tmp_path / "b"
    out.mkdir()
    cls = np.zeros((20, 20), np.uint8)
    cls[5:9, 5:9] = 1
    path = write_raster(out / "building_mask.tif", cls, LONDON, dtype="uint8")
    result = SimpleNamespace(success=True, pipeline="building_footprints", output_path=path, stats={"building_coverage": 0.04},
                             metadata={"date": "2024-06"}, duration_seconds=1.0)
    d = generate_deliverables(result, LONDON, out, report_formats="docx")
    assert "accuracy is not assessed" in docx_text(d.reports["docx"])
    assert not [w for w in d.warnings if "expected output" in w], "the dataset layers are optional in network mode"


def test_road_network_report(tmp_path, monkeypatch):
    bbox = (0.45, 5.90, 0.47, 5.92)
    lines = [LineString([(0.455, 5.91), (0.465, 5.91)]), LineString([(0.455, 5.915), (0.46, 5.915)])]
    osm = gpd.GeoDataFrame({"highway": ["primary", "residential"], "name": ["Main", None], "surface": [None, None],
                            "oneway": [None, None], "osm_id": [1, 2]}, geometry=lines, crs="EPSG:4326")
    monkeypatch.setattr(fp, "osm_features", lambda b, kind, **k: osm)
    from pygeovision.ai.pipelines.vector_products import RoadNetworkPipeline
    r = RoadNetworkPipeline(None).run(bbox, tmp_path / "out")
    d = build(r, bbox, tmp_path / "out", vector_formats="gpkg,shp")
    text = docx_text(d.reports["docx"])
    assert "Road length by OpenStreetMap highway class" in text and "primary" in text and "OpenStreetMap" in text
    roads = gpd.read_file(d.vectors["gpkg"][0], layer="roads")
    assert len(roads) == 2 and roads.geom_type.eq("LineString").all()
    shp = gpd.read_file(next(p for p in d.vectors["shp"] if p.name.endswith("_roads.shp")))
    assert len(shp) == 2 and {"length_m", "highway"} <= set(shp.columns)


def test_water_extent_timeseries_report(tmp_path, monkeypatch):
    def field(c):
        a = np.zeros((N, N), "float32")
        a[:, :c] = 1
        return a
    scenes = {"2024-01": field(10), "2024-02": field(20), "2024-03": field(15)}
    src = {k: tif(tmp_path, f"w{k}", mndwi_bands(v)) for k, v in scenes.items()}
    patch_downloads(monkeypatch, src)
    from pygeovision.ai.pipelines.timeseries import WaterExtentTimeseriesPipeline
    r = WaterExtentTimeseriesPipeline(None).run(LONDON, tmp_path / "out", dates=list(scenes))
    d = build(r, LONDON, tmp_path / "out")
    text = docx_text(d.reports["docx"])
    assert "Water area by date" in text and "2024-02" in text
    assert len(d.charts) >= 2 and json.loads(d.summary_json.read_text())["period"] == "2024-01 to 2024-03"


def test_vegetation_anomaly_report(tmp_path, monkeypatch):
    s = {f"{2019 + i}-06": tif(tmp_path, f"b{i}", red_nir(ndvi_field(v, v))) for i, v in enumerate([0.5, 0.6, 0.5, 0.6, 0.5])}
    s["2024-06"] = tif(tmp_path, "t", red_nir(ndvi_field(0.2, 0.55)))
    patch_downloads(monkeypatch, s)
    from pygeovision.ai.pipelines.timeseries import VegetationAnomalyPipeline
    r = VegetationAnomalyPipeline(None).run(LONDON, tmp_path / "out", date="2024-06")
    d = build(r, LONDON, tmp_path / "out")
    text = docx_text(d.reports["docx"])
    assert "Much below normal" in text and "Vegetation Condition Index" in text


@pytest.mark.parametrize("pipeline,fname", [("road_extraction", "road_mask.tif"), ("landslide_detection", "landslide_map.tif")])
def test_weight_gated_pipelines_have_honest_reports(tmp_path, pipeline, fname):
    out = tmp_path / pipeline
    out.mkdir()
    cls = np.zeros((20, 20), np.uint8)
    cls[5:9, 5:9] = 1
    path = write_raster(out / fname, cls, LONDON, dtype="uint8")
    stats = {"road_coverage": 0.04} if pipeline == "road_extraction" else {"changed_fraction": 0.04}
    meta = {"date": "2024-06"} if pipeline == "road_extraction" else {"date_before": "2024-01", "date_after": "2024-02"}
    result = SimpleNamespace(success=True, pipeline=pipeline, output_path=path, stats=stats, metadata=meta, duration_seconds=1.0)
    d = generate_deliverables(result, LONDON, out, report_formats="docx")
    assert "accuracy is not assessed" in docx_text(d.reports["docx"])
    assert not [w for w in d.warnings if "expected output" in w]


def test_all_new_specs_exist_and_name_real_files():
    from pygeovision.ai.pipelines import list_pipelines
    assert set(SPECS) == set(list_pipelines())
    for name in ("landcover_change", "urban_expansion", "forest_change", "surface_water_dynamics", "evapotranspiration", "active_fire",
                 "terrain_analysis", "drainage_analysis", "vegetation_anomaly", "water_extent_timeseries", "road_network", "sar_flood_mapping", "multisensor_flood_mapping"):
        s = SPECS[name]
        assert s.method and s.limitations and s.references and s.layers
