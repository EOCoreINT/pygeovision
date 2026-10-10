"""Map, vector layers and reports (PDF / Word / Excel) for every pipeline run.

The numbers in a report are measured from the output rasters, so these tests build
rasters whose answer is known by construction (a rectangle of known size at a known
latitude) and check the areas, the files, and that one failing part does not lose the rest.
"""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

pytest.importorskip("rasterio")
pytest.importorskip("geopandas")
pytest.importorskip("matplotlib")
pytest.importorskip("reportlab")
pytest.importorskip("docx")
pytest.importorskip("openpyxl")

from pipeline_synth import BBOX, write_raster  # noqa: E402  (tests/ is on sys.path)

from pygeovision.reporting import (  # noqa: E402
    REPORT_FORMATS, SPECS, generate_deliverables, regenerate,
)


def geod_area_ha(bounds):
    from pyproj import Geod
    w, s, e, n = bounds
    a, _ = Geod(ellps="WGS84").polygon_area_perimeter([w, e, e, w], [s, s, n, n])
    return abs(a) / 1e4


@pytest.fixture
def flood_run(tmp_path):
    """A flood_mapping output: 40x40 px, a 10x20 block 'flooded', a 5x40 strip 'permanent', one nodata column."""
    out = tmp_path / "flood"
    out.mkdir()
    cls = np.zeros((40, 40), np.uint8)
    cls[5:15, 5:25] = 1
    cls[30:35, :] = 2
    cls[:, 39] = 255
    path = write_raster(out / "flood_extent.tif", cls, BBOX, dtype="uint8", nodata=255)
    stats = {"flooded_area_ha": 1.0, "permanent_water_area_ha": 1.0, "flooded_fraction": 0.1, "valid_fraction": 0.975, "mndwi_threshold": 0.0}
    result = SimpleNamespace(success=True, pipeline="flood_mapping", output_path=path, stats=stats,
                             metadata={"date_before": "2024-05", "date_after": "2024-08"}, duration_seconds=1.5)
    return result, out, cls


# ---------------------------------------------------------------------------
# Measurement
# ---------------------------------------------------------------------------
class TestAreas:
    def test_row_areas_match_geodesic_polygon(self, tmp_path):
        import rasterio
        from pygeovision.reporting.geo import row_pixel_area_ha
        path = write_raster(tmp_path / "a.tif", np.zeros((50, 50), np.uint8), BBOX, dtype="uint8")
        with rasterio.open(path) as src:
            rows = row_pixel_area_ha(src)
            total = float(rows.sum() * src.width)
        assert total == pytest.approx(geod_area_ha(BBOX), rel=2e-4)
        assert rows[0] < rows[-1], "pixels get larger towards the equator (south) for a northern-hemisphere box"

    def test_class_table_counts_and_nodata(self, flood_run):
        from pygeovision.reporting.geo import class_table
        from pygeovision.reporting.palettes import ClassDef
        result, out, cls = flood_run
        defs = {0: ClassDef(0, "Dry", "#ccc"), 1: ClassDef(1, "Flooded", "#e00"), 2: ClassDef(2, "Permanent", "#00e")}
        t = class_table(result.output_path, 1, defs)
        by = {r["value"]: r for r in t.rows}
        assert by[1]["pixels"] == 200 and by[2]["pixels"] == 195
        assert t.valid_pixels == 40 * 39 and t.total_pixels == 1600, "the 255 column is nodata, not a class"
        assert sum(r["share"] for r in t.rows) == pytest.approx(1.0)
        assert t.total_area_ha == pytest.approx(geod_area_ha(BBOX), rel=2e-4)
        # 200 of 1600 pixels of the full raster (the nodata column cuts 5 px from the strip)
        assert by[1]["area_ha"] == pytest.approx(geod_area_ha(BBOX) * 200 / 1600, rel=0.01)

    @pytest.mark.parametrize("kind,kw", [
        ("fixed", dict(edges=(0.25, 0.5), names=("a", "b", "c"), colors=("#111", "#222", "#333"))),
        ("equal", dict(n=4, unit="C")),
        ("quantile", dict(n=3)),
        ("std", dict(k=1.0, names=("lo", "mid", "hi"), colors=("#111", "#222", "#333"))),
    ])
    def test_binning_covers_every_valid_pixel_and_keeps_nan_invalid(self, kind, kw):
        from pygeovision.reporting.geo import apply_binning
        from pygeovision.reporting.specs import Binning
        rng = np.random.default_rng(1)
        arr = rng.random((30, 30)).astype(np.float32)
        arr[0, :5] = np.nan
        cls, defs = apply_binning(arr, Binning(kind, **kw))
        assert (cls[0, :5] == 255).all() and (cls[1:] != 255).all()
        assert {d.value for d in defs} >= set(np.unique(cls[cls != 255]).tolist())

    def test_equal_bin_names_stay_distinct_for_a_tiny_range(self):
        from pygeovision.reporting.geo import apply_binning
        from pygeovision.reporting.specs import Binning
        arr = (-124.1 + np.linspace(0, 0.002, 400)).reshape(20, 20).astype(np.float32)
        _, defs = apply_binning(arr, Binning("equal", n=5, unit="C"))
        assert len({d.name for d in defs}) == 5


# ---------------------------------------------------------------------------
# Vectors
# ---------------------------------------------------------------------------
class TestVectors:
    def _poly(self, arr, **kw):
        from rasterio.transform import from_origin
        from pygeovision.reporting.palettes import ClassDef
        from pygeovision.reporting.vectors import polygonise
        h, w = arr.shape
        tr = from_origin(BBOX[0], BBOX[3], (BBOX[2] - BBOX[0]) / w, (BBOX[3] - BBOX[1]) / h)
        defs = [ClassDef(0, "bg", "#ccc"), ClassDef(1, "one", "#e00"), ClassDef(2, "two", "#00e")]
        args = dict(keep=(1, 2), name="x", source_layer="x", aoi_area_ha=geod_area_ha(BBOX), pipeline="p", period="2024")
        args.update(kw)
        return polygonise(arr, tr, "EPSG:4326", defs, **args)

    def test_rectangle_area_matches_geodesic_area(self):
        arr = np.zeros((40, 40), np.float32)
        arr[10:20, 10:30] = 1
        lyr = self._poly(arr)
        assert lyr.n_features == 1
        row = lyr.gdf.iloc[0]
        w = (BBOX[2] - BBOX[0]) / 40
        h = (BBOX[3] - BBOX[1]) / 40
        box = (BBOX[0] + 10 * w, BBOX[3] - 20 * h, BBOX[0] + 30 * w, BBOX[3] - 10 * h)
        assert row["area_ha"] == pytest.approx(geod_area_ha(box), rel=0.003)
        assert row["class_name"] == "one" and row["color_hex"] == "#e00"
        assert row["pct_aoi"] == pytest.approx(row["area_ha"] / geod_area_ha(BBOX) * 100, rel=1e-6)
        assert str(lyr.gdf.crs).upper() == "EPSG:4326"

    def test_invalid_pixels_and_skipped_classes_are_not_polygons(self):
        arr = np.zeros((20, 20), np.float32)
        arr[2:6, 2:6] = 1
        arr[10:14, 10:14] = 2
        arr[2:6, 12:16] = np.nan
        lyr = self._poly(arr, keep=None, skip=(0, 2))
        assert set(lyr.gdf["class_id"]) == {1}
        assert lyr.n_features == 1

    def test_too_many_patches_are_generalised_and_the_layer_says_so(self):
        rng = np.random.default_rng(0)
        arr = (rng.random((80, 80)) > 0.5).astype(np.float32)          # salt-and-pepper speckle
        raw = self._poly(arr, max_polygons=10_000_000)
        lyr = self._poly(arr, max_polygons=200)
        assert raw.n_features > 200
        assert lyr.n_features <= 200 and lyr.sieve_pixels >= 4 and "merged" in lyr.note

    def test_all_formats_are_written_and_attributes_survive_the_shapefile(self, tmp_path):
        import geopandas as gpd
        from pygeovision.reporting.vectors import aoi_frame, write_vectors
        arr = np.zeros((40, 40), np.float32)
        arr[10:20, 10:30] = 1
        arr[25:30, 5:9] = 2
        lyr = self._poly(arr, name="flood_extent")
        aoi = aoi_frame(BBOX, {"pipeline": "p", "area_km2": 19.3})
        written, warns = write_vectors([lyr], aoi, tmp_path / "v", "p", ["gpkg", "geojson", "shp"])
        assert not warns, warns
        layers = gpd.list_layers(written["gpkg"][0])["name"].tolist()
        assert {"aoi", "flood_extent", "flood_extent_classes"} <= set(layers)
        gj = gpd.read_file(next(p for p in written["geojson"] if "flood_extent" in p.name))
        assert len(gj) == 2 and gj.crs.to_epsg() == 4326
        shp = gpd.read_file(next(p for p in written["shp"] if "flood_extent" in p.name))
        assert {"class_id", "class_name", "area_ha", "pct_aoi", "color_hex"} <= set(shp.columns), "no 10-character truncation"
        assert shp["area_ha"].sum() == pytest.approx(lyr.gdf["area_ha"].sum(), rel=1e-6)

    def test_unknown_format_is_reported_not_raised(self, tmp_path):
        from pygeovision.reporting.vectors import aoi_frame, write_vectors
        _, warns = write_vectors([], aoi_frame(BBOX, {"a": 1}), tmp_path, "p", ["gpkg", "dxf"])
        assert any("dxf" in w for w in warns)


# ---------------------------------------------------------------------------
# End to end
# ---------------------------------------------------------------------------
class TestDeliverables:
    def test_every_pipeline_has_a_report_spec(self):
        from pygeovision.ai.pipelines import list_pipelines
        assert set(SPECS) == set(list_pipelines()), "a new pipeline needs a ReportSpec in reporting/specs.py"

    def test_spec_text_fields_are_tuples_not_strings(self):
        for name, s in SPECS.items():
            for f in ("method", "reading", "limitations", "references"):
                assert not isinstance(getattr(s, f), str), f"{name}.{f} is a string; a stray missing comma?"

    def test_flood_run_produces_everything_and_the_numbers_agree(self, flood_run):
        result, out, cls = flood_run
        d = generate_deliverables(result, BBOX, out, parameters={"date_before": "2024-05"}, vector_formats=["gpkg", "geojson", "shp"])
        assert not d.warnings, d.warnings
        assert set(d.reports) == set(REPORT_FORMATS)
        assert all(p.exists() and p.stat().st_size > 100 for p in d.files())
        assert d.maps and d.charts and d.summary_json.exists()

        summary = json.loads(d.summary_json.read_text())
        rows = {r["name"]: r for r in summary["layers"][0]["classes"]}
        px = geod_area_ha(BBOX) / 1600
        assert rows["Flooded"]["area_ha"] == pytest.approx(200 * px, rel=0.01)
        assert rows["Permanent water"]["area_ha"] == pytest.approx(195 * px, rel=0.01)

        import geopandas as gpd
        gdf = gpd.read_file(d.vectors["gpkg"][0], layer="flood_extent")
        assert gdf.groupby("class_name")["area_ha"].sum()["Flooded"] == pytest.approx(rows["Flooded"]["area_ha"], rel=0.01)

    def test_pdf_docx_xlsx_contain_the_results(self, flood_run):
        import openpyxl
        from docx import Document
        pypdf = pytest.importorskip("pypdf")
        result, out, _ = flood_run
        d = generate_deliverables(result, BBOX, out, parameters={})
        reader = pypdf.PdfReader(str(d.reports["pdf"]))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        assert len(reader.pages) >= 3
        assert "Flood Extent Mapping" in text and "Flooded" in text and "Permanent water" in text
        assert "Executive summary" in text and "Limitations" in text and "n/a" not in text.split("Method")[0]

        doc = Document(str(d.reports["docx"]))
        body = "\n".join(p.text for p in doc.paragraphs)
        cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        assert "Executive summary" in body and "Flooded" in cells and "Area (ha)" in cells
        assert len(doc.inline_shapes) >= 2, "map and chart are embedded"

        wb = openpyxl.load_workbook(d.reports["xlsx"])
        assert {"Summary", "Metrics", "Classes", "Files"} <= set(wb.sheetnames)
        ws = wb["Classes"]
        formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        assert any("SUM" in f for f in formulas), "totals and shares are live formulas"
        assert ws._charts, "native Excel chart"

    def test_one_failing_writer_does_not_lose_the_others(self, flood_run, monkeypatch):
        import pygeovision.reporting.pdf as pdf_mod
        monkeypatch.setattr(pdf_mod, "write_pdf", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("disk on fire")))
        result, out, _ = flood_run
        d = generate_deliverables(result, BBOX, out)
        assert "pdf" not in d.reports and {"docx", "xlsx"} <= set(d.reports)
        assert any("disk on fire" in w for w in d.warnings)
        assert d.maps and d.vectors

    def test_missing_output_raster_is_reported_and_stats_still_delivered(self, flood_run):
        result, out, _ = flood_run
        result.output_path.unlink()
        d = generate_deliverables(result, BBOX, out)
        assert any("flood_extent.tif" in w for w in d.warnings)
        assert {"pdf", "docx", "xlsx"} <= set(d.reports)

    def test_failed_run_is_refused(self, tmp_path):
        with pytest.raises(ValueError, match="successful"):
            generate_deliverables(SimpleNamespace(success=False), BBOX, tmp_path)

    def test_formats_can_be_chosen_and_unknown_ones_warn(self, flood_run):
        result, out, _ = flood_run
        d = generate_deliverables(result, BBOX, out, report_formats="pdf,odt", vector_formats="none")
        assert set(d.reports) == {"pdf"} and not d.vectors
        assert any("odt" in w for w in d.warnings)

    def test_class_names_label_a_trained_network(self, tmp_path):
        out = tmp_path / "crop"
        out.mkdir()
        cls = np.zeros((30, 30), np.uint8)
        cls[:10] = 1
        cls[10:20] = 2
        path = write_raster(out / "crop_map.tif", cls, BBOX, dtype="uint8")
        result = SimpleNamespace(success=True, pipeline="crop_monitoring", output_path=path, stats={},
                                 metadata={"date": "2023-06"}, duration_seconds=1.0)
        d = generate_deliverables(result, BBOX, out, class_names=["Background", "Maize", "Wheat"], report_formats="pdf")
        import geopandas as gpd
        gdf = gpd.read_file(d.vectors["gpkg"][0], layer="crop_fields")
        assert set(gdf["class_name"]) == {"Maize", "Wheat"}, "class 0 (background) is not polygonised"
        text = json.loads(d.summary_json.read_text())["layers"][0]["classes"]
        assert {r["name"] for r in text} == {"Background", "Maize", "Wheat"}

    def test_trained_network_report_says_accuracy_is_not_assessed(self, tmp_path):
        out = tmp_path / "b"
        out.mkdir()
        cls = np.zeros((20, 20), np.uint8)
        cls[5:9, 5:9] = 1
        path = write_raster(out / "building_mask.tif", cls, BBOX, dtype="uint8")
        result = SimpleNamespace(success=True, pipeline="building_footprints", output_path=path,
                                 stats={"building_coverage": 0.04}, metadata={"date": "2024-06"}, duration_seconds=1.0)
        d = generate_deliverables(result, BBOX, out, report_formats="docx")
        from docx import Document
        text = "\n".join(c.text for t in Document(str(d.reports["docx"])).tables for r in t.rows for c in r.cells)
        assert "accuracy is not assessed" in text

    def test_low_coverage_is_a_visible_warning(self, flood_run):
        result, out, _ = flood_run
        result.stats = {**result.stats, "aoi_coverage_fraction": 0.4}
        d = generate_deliverables(result, BBOX, out, report_formats="docx", vector_formats="none")
        from docx import Document
        text = "\n".join(c.text for t in Document(str(d.reports["docx"])).tables for r in t.rows for c in r.cells)
        assert "covers only 40 %" in text

    def test_regenerate_from_run_summary(self, flood_run):
        result, out, _ = flood_run
        generate_deliverables(result, BBOX, out, report_formats="pdf")
        d = regenerate(out, report_formats="docx", vector_formats="shp")
        assert set(d.reports) == {"docx"} and "shp" in d.vectors

    def test_regenerate_without_summary_says_what_is_wrong(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="run_summary.json"):
            regenerate(tmp_path)

    def test_large_negative_zero_is_not_printed(self):
        from pygeovision.reporting.fmt import fmt_value
        assert fmt_value(-0.0004, "num2") == "0.00"
        assert fmt_value(-0.4, "num1") == "-0.4"


# ---------------------------------------------------------------------------
# Wiring: facade and CLI
# ---------------------------------------------------------------------------
@pytest.fixture
def offline_flood(tmp_path, monkeypatch):
    """flood_mapping with the download replaced by two synthetic scenes."""
    from pygeovision.ai.pipelines import BasePipeline
    rng = np.random.default_rng(0)
    a = np.clip(rng.normal(0.2, 0.1, (2, 40, 40)), 0.02, 0.9).astype("float32")
    b = np.clip(a + rng.normal(0, 0.1, a.shape), 0.02, 0.9).astype("float32")
    pa, pb = write_raster(tmp_path / "a.tif", a, BBOX), write_raster(tmp_path / "b.tif", b, BBOX)
    monkeypatch.setattr(BasePipeline, "_search_and_download_pair", lambda self, *a_, **k: (pa, pb))
    return tmp_path


class TestWiring:
    def test_facade_creates_deliverables_by_default(self, offline_flood, monkeypatch):
        monkeypatch.delenv("PYGEOVISION_REPORT", raising=False)
        from pygeovision import PyGeoVision
        r = PyGeoVision().pipeline("flood_mapping", bbox=BBOX, output_dir=str(offline_flood / "o"),
                                   date_before="2024-05", date_after="2024-08", report_formats="pdf", vector_formats="gpkg")
        assert r.success and r.deliverables is not None
        assert r.deliverables.reports["pdf"].exists() and r.deliverables.vectors["gpkg"][0].exists()
        params = json.loads((offline_flood / "o" / "run_summary.json").read_text())["parameters"]
        assert params["date_before"] == "2024-05" and "report_formats" not in params

    def test_report_false_and_env_switch_it_off(self, offline_flood, monkeypatch):
        from pygeovision import PyGeoVision
        kw = dict(bbox=BBOX, date_before="2024-05", date_after="2024-08")
        r = PyGeoVision().pipeline("flood_mapping", output_dir=str(offline_flood / "o1"), report=False, **kw)
        assert r.success and r.deliverables is None and not (offline_flood / "o1" / "deliverables").exists()
        monkeypatch.setenv("PYGEOVISION_REPORT", "0")
        r = PyGeoVision().pipeline("flood_mapping", output_dir=str(offline_flood / "o2"), **kw)
        assert r.deliverables is None and not (offline_flood / "o2" / "deliverables").exists()

    def test_a_reporting_crash_never_fails_the_analysis(self, offline_flood, monkeypatch):
        import pygeovision.reporting as rep
        monkeypatch.setattr(rep, "generate_deliverables", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
        from pygeovision import PyGeoVision
        r = PyGeoVision().pipeline("flood_mapping", bbox=BBOX, output_dir=str(offline_flood / "o"),
                                   date_before="2024-05", date_after="2024-08")
        assert r.success and r.deliverables is None and "boom" in r.metadata["deliverables_error"]
        assert (offline_flood / "o" / "flood_extent.tif").exists()

    def test_cli_channel_options_and_report_command(self, offline_flood, monkeypatch):
        monkeypatch.delenv("PYGEOVISION_REPORT", raising=False)
        from click.testing import CliRunner
        from pygeovision.cli.main import cli
        runner = CliRunner()
        base = ["channel", "flood_mapping", "--bbox", *map(str, BBOX), "--date-before", "2024-05", "--date-after", "2024-08"]
        out = offline_flood / "c1"
        res = runner.invoke(cli, [*base, "-o", str(out), "--report-formats", "docx", "--vector-formats", "geojson"])
        assert res.exit_code == 0, res.output
        assert "Deliverables:" in res.output and "Word report" in res.output and "Map" in res.output
        assert (out / "deliverables" / "flood_mapping_report.docx").exists()
        assert not (out / "deliverables" / "flood_mapping_report.pdf").exists()

        res = runner.invoke(cli, ["report", str(out), "--formats", "pdf", "--vector-formats", "shp"])
        assert res.exit_code == 0, res.output
        assert (out / "deliverables" / "flood_mapping_report.pdf").exists()
        assert list((out / "deliverables" / "vectors" / "shapefile").glob("*.shp"))

        out2 = offline_flood / "c2"
        res = runner.invoke(cli, [*base, "-o", str(out2), "--no-report"])
        assert res.exit_code == 0 and "Deliverables" not in res.output
        assert not (out2 / "deliverables").exists()

    def test_report_command_on_a_folder_that_is_not_a_run(self, tmp_path):
        from click.testing import CliRunner
        from pygeovision.cli.main import cli
        res = CliRunner().invoke(cli, ["report", str(tmp_path)])
        assert res.exit_code == 1 and "run_summary.json" in res.output
